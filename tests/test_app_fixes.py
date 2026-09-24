"""Тесты трёх падений в app.py (запускать из docx_converter/)."""
import io
import socket
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Pt, RGBColor
from pypdf import PdfReader, PdfWriter

import app
from convert import STYLES, convert_md_to_docx
from ocr.mineru_provider import MineruAuthError, MineruError
from ocr_fixtures import require_fixture
from pdf_core import PageInfo


class _FakeUpload:
    """Минимальный стенд вместо streamlit UploadedFile."""

    def __init__(self, name: str, data: bytes = b"fake"):
        self.name = name
        self._data = data

    def getvalue(self) -> bytes:
        return self._data


def test_page_range_ignored_for_non_pdf(monkeypatch):
    seen = []

    def fake_convert(path, page_range=None):
        seen.append(page_range)
        return "# Документ"

    monkeypatch.setattr(app, "convert_with_markitdown", fake_convert)

    result = app._convert_uploaded_file(_FakeUpload("отчёт.docx"), "1-3")

    assert seen == [None]
    assert result["error"] is None
    assert result["markdown"] == "# Документ"
    assert result["page_range"] == "all"


class _RaisingSecrets:
    """st.secrets на машине без .streamlit/secrets.toml."""

    def __contains__(self, key):
        raise FileNotFoundError("No secrets files found")


def test_drive_unavailable_without_secrets_file(monkeypatch):
    monkeypatch.setattr(app.st, "secrets", _RaisingSecrets())

    assert app._drive_secrets_available() is False


def test_drive_available_with_service_account(monkeypatch):
    monkeypatch.setattr(app.st, "secrets", {"gcp_service_account": {}})

    assert app._drive_secrets_available() is True


BOM_FIXTURE = Path(__file__).resolve().parents[1] / "test_formatting_bom.md"


def test_decode_md_upload_strips_bom():
    md_text = app._decode_md_upload(BOM_FIXTURE.read_bytes())

    assert "﻿" not in md_text
    assert md_text.startswith("# ")


def test_decode_md_upload_falls_back_to_cp1251():
    assert app._decode_md_upload("Тест кириллицы".encode("cp1251")) == "Тест кириллицы"


def test_bom_fixture_first_block_renders_as_h1(tmp_path):
    """BOM не должен превращать H1 в обычный абзац."""
    md_text = app._decode_md_upload(BOM_FIXTURE.read_bytes())
    out = tmp_path / "bom.docx"
    convert_md_to_docx(md_text, str(out))

    first = Document(str(out)).paragraphs[0]
    run = first.runs[0]
    assert not first.text.startswith("#")
    assert run.font.name == "PT Sans Narrow"
    assert run.font.size == Pt(18)
    assert run.font.color.rgb == RGBColor.from_string("015198")


def test_every_doc_type_has_a_known_style():
    """ПРАВКА #59: опечатка в "style" отдала бы клиенту письмо в стиле ПЗ —
    ловим её здесь, а не в готовом документе."""
    for name, cfg in app.DOC_TYPES.items():
        assert "style" in cfg, f"у типа {name} нет ключа style"
        assert cfg["style"] in STYLES, f"у типа {name} неизвестный style: {cfg['style']}"


def test_doc_types_point_at_distinct_templates():
    """ПРАВКА #54: раньше оба типа грузили один и тот же file_id."""
    drive_ids = [cfg["drive_id"] for cfg in app.DOC_TYPES.values()]

    assert len(set(drive_ids)) == len(drive_ids)
    assert all(drive_ids)


def test_get_template_downloads_selected_drive_id(monkeypatch):
    seen = []

    def fake_download(file_id):
        seen.append(file_id)
        return "/tmp/template.docx"

    monkeypatch.setattr(app, "download_template_from_drive", fake_download)

    path = app.get_template(use_drive=True,
                            drive_file_id="1_E7eI5PgMD50MEI8RNl8xoiWmhUsOUap",
                            local_path=None)

    assert seen == ["1_E7eI5PgMD50MEI8RNl8xoiWmhUsOUap"]
    assert path == "/tmp/template.docx"


# --- ПРАВКА #73: MinerU в режиме «Файлы -> Markdown» -------------------------

@pytest.fixture
def no_network(monkeypatch):
    """ПРАВКА #85: как в tests/test_ocr_golden.py, но только для тестов блока #73/#85."""
    def refuse(*args, **kwargs):
        raise AssertionError("тест полез в сеть")
    monkeypatch.setattr(socket, "socket", refuse)


_offline = pytest.mark.usefixtures("no_network")


def _blank_pdf(pages: int) -> bytes:
    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def _make_zip(full_md: str) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("full.md", full_md)
    return buffer.getvalue()


def _fake_mineru(monkeypatch, tmp_path, **texts):
    """Фейковый провайдер вместо сети + кэш в tmp_path. Возвращает журнал вызовов."""
    calls: list[str] = []

    class FakeProvider:
        def __init__(self, mv):
            self.mv = mv

        def page_infos(self, pdf_bytes, page_range=None):
            return [PageInfo(1, False, True)]

        def fetch_raw_zip(self, pdf_bytes, page_range=None):
            calls.append(self.mv)
            return _make_zip(texts[self.mv])

    monkeypatch.setattr(app, "_mineru_provider_factory",
                        lambda api_key, status: FakeProvider)
    monkeypatch.setattr(app, "_OCR_CACHE_ROOT", tmp_path / "cache")
    return calls


def test_ocr_mode_options_hide_ocrmypdf_without_binaries():
    """ПРАВКА #73: на Streamlit Cloud бинарников нет — OCRmyPDF не предлагается."""
    assert app._ocr_mode_options(False) == ["off", "mineru"]
    assert app._ocr_mode_options(True) == ["off", "auto", "mineru"]


def test_mineru_api_key_prefers_secrets(monkeypatch):
    monkeypatch.setattr(app.st, "secrets", {"MINERU_API_KEY": "from-secrets"})
    monkeypatch.setenv("MINERU_API_KEY", "from-env")

    assert app._mineru_api_key() == "from-secrets"


class _NoSecretsFile:
    def __getitem__(self, key):
        raise FileNotFoundError("No secrets files found")


def test_mineru_api_key_falls_back_to_env(monkeypatch):
    """ПРАВКА #73: нет secrets.toml или нет ключа в нём — берём переменную окружения."""
    monkeypatch.setenv("MINERU_API_KEY", "from-env")

    monkeypatch.setattr(app.st, "secrets", _NoSecretsFile())
    assert app._mineru_api_key() == "from-env"

    monkeypatch.setattr(app.st, "secrets", {"gemini": {}})
    assert app._mineru_api_key() == "from-env"


def test_mineru_api_key_absent(monkeypatch):
    monkeypatch.setattr(app.st, "secrets", {})
    monkeypatch.delenv("MINERU_API_KEY", raising=False)

    assert app._mineru_api_key() is None


@_offline
def test_mineru_mode_runs_ingest_with_fake_provider(monkeypatch, tmp_path):
    """ПРАВКА #73/#85: UI зовёт ingest, в результате — markdown и отчёт."""
    calls = _fake_mineru(monkeypatch, tmp_path, vlm="# Договор\n\nТекст договора.")

    result = app._convert_uploaded_file(_FakeUpload("скан.pdf", _blank_pdf(1)), None,
                                        ocr_mode="mineru")

    assert result["error"] is None
    assert "Текст договора." in result["markdown"]
    assert result["report"]["provider"] == "mineru"
    assert result["report"]["verified"] is False
    assert calls == ["vlm"]
    assert result["route"] == "scan"


@_offline
def test_mineru_mode_verify_runs_second_model(monkeypatch, tmp_path):
    calls = _fake_mineru(monkeypatch, tmp_path,
                         vlm="Цена 100 рублей за штуку товара.",
                         pipeline="Цена 700 рублей за штуку товара.")

    result = app._convert_uploaded_file(_FakeUpload("скан.pdf", _blank_pdf(1)), None,
                                        ocr_mode="mineru", verify=True)

    assert result["error"] is None
    assert calls == ["vlm", "pipeline"]
    assert result["report"]["verified"] is True
    _, low_conf = app._split_findings(result["report"])
    assert low_conf, "расхождение 100/700 должно стать находкой low_confidence"


@_offline
def test_mineru_auth_error_gets_key_hint(monkeypatch, tmp_path):
    """ПРАВКА #73: ошибка ключа — текстом с подсказкой, не трейсбеком."""
    def factory(api_key, status):
        def build(model_version):
            raise MineruAuthError("A0202: token invalid")
        return build

    monkeypatch.setattr(app, "_mineru_provider_factory", factory)
    monkeypatch.setattr(app, "_OCR_CACHE_ROOT", tmp_path / "cache")

    result = app._convert_uploaded_file(_FakeUpload("скан.pdf", _blank_pdf(1)), None,
                                        ocr_mode="mineru")

    assert "A0202" in result["error"]
    assert "MINERU_API_KEY" in result["error"]
    assert result["report"] is None


@_offline
def test_mineru_error_shown_as_text_without_hint(monkeypatch, tmp_path):
    def factory(api_key, status):
        def build(model_version):
            raise MineruError("распознавание не удалось: boom")
        return build

    monkeypatch.setattr(app, "_mineru_provider_factory", factory)
    monkeypatch.setattr(app, "_OCR_CACHE_ROOT", tmp_path / "cache")

    result = app._convert_uploaded_file(_FakeUpload("скан.pdf", _blank_pdf(1)), None,
                                        ocr_mode="mineru")

    assert result["error"] == "распознавание не удалось: boom"


def test_pdf_page_subset_cuts_selected_pages():
    """ПРАВКА #73: run_pipeline диапазона не знает — страницы вырезает app.py."""
    pdf = _blank_pdf(3)

    subset = app._pdf_page_subset(pdf, "2-3")

    assert len(PdfReader(io.BytesIO(subset)).pages) == 2
    assert app._pdf_page_subset(pdf, None) is pdf


def test_pdf_page_subset_rejects_page_outside_pdf():
    with pytest.raises(ValueError, match="всего 3"):
        app._pdf_page_subset(_blank_pdf(3), "2-5")


def test_split_findings_separates_low_confidence():
    report = {"findings": [{"rule": "inn_checksum"}, {"rule": "low_confidence"},
                           {"rule": "gost_format"}]}

    main, low_conf = app._split_findings(report)

    assert [f["rule"] for f in main] == ["inn_checksum", "gost_format"]
    assert [f["rule"] for f in low_conf] == ["low_confidence"]


# --- ПРАВКА #85: UI на едином входе ocr.ingest -------------------------------

REPORT_KEYS = ["schema_version", "source", "sha256", "provider", "model_version",
               "cache_hit", "verified", "created_at", "summary", "findings"]
RESULT_KEYS = {"filename", "download_name", "file_type", "page_range", "ocr_status",
               "markdown", "report", "route", "error"}


def _docx_bytes() -> bytes:
    doc = Document()
    doc.add_heading("Договор поставки", level=1)
    doc.add_paragraph("Поставщик обязуется передать товар.")
    table = doc.add_table(rows=2, cols=2)
    for r, row in enumerate((("Товар", "Цена"), ("Весы", "100"))):
        for c, text in enumerate(row):
            table.cell(r, c).text = text
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _xlsx_bytes() -> bytes:
    from openpyxl import Workbook

    book = Workbook()
    for row in (("Позиция", "Сумма"), (1, 100), (2, 250)):
        book.active.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


@_offline
@pytest.mark.parametrize("name, data", [("док.docx", _docx_bytes()), ("табл.xlsx", _xlsx_bytes())],
                         ids=["docx", "xlsx"])   # без ids байты файла уходят в id теста: PYTEST_CURRENT_TEST > 32767 на Windows
def test_mineru_mode_office_goes_through_ingest(monkeypatch, tmp_path, name, data):
    calls = _fake_mineru(monkeypatch, tmp_path, vlm="не должен понадобиться")
    result = app._convert_uploaded_file(_FakeUpload(name, data), "1-3", ocr_mode="mineru", verify=True)
    assert result["error"] is None and result["markdown"].strip()
    assert result["route"] == "office" and result["page_range"] == "all"
    assert result["report"]["provider"] == "markitdown"
    assert result["report"]["verified"] is False               # verify не передан, не ValueError
    assert list(result["report"]) == REPORT_KEYS                # те же десять ключей, что у PDF
    assert calls == []                                          # провайдер не создавался


@_offline
def test_mineru_mode_text_pdf_stays_local(monkeypatch, tmp_path):
    pdf = require_fixture("textpdf1.pdf").read_bytes()          # стр. 1-2: слой есть, таблиц нет (спека 09)
    calls = _fake_mineru(monkeypatch, tmp_path, vlm="x")
    result = app._convert_uploaded_file(_FakeUpload("т.pdf", pdf), "1-2", ocr_mode="mineru")
    assert result["route"] == "text" and calls == []            # маршрут — по вырезке
    assert result["report"]["provider"] == "markitdown"


@_offline
def test_mineru_mode_pptx_keeps_plain_markitdown(monkeypatch):
    monkeypatch.setattr(app, "convert_with_markitdown", lambda path, page_range=None: "# Слайд")
    result = app._convert_uploaded_file(_FakeUpload("през.pptx"), None, ocr_mode="mineru")
    assert result["markdown"] == "# Слайд" and result["report"] is None and result["route"] is None


@_offline
def test_office_outside_mineru_mode_unchanged(monkeypatch):     # off: отчёта нет, как было
    monkeypatch.setattr(app, "convert_with_markitdown", lambda path, page_range=None: "# Д")
    result = app._convert_uploaded_file(_FakeUpload("д.docx"), None, ocr_mode="off")
    assert result["report"] is None and result["route"] is None


@_offline
def test_result_keys_closed(monkeypatch, tmp_path):
    _fake_mineru(monkeypatch, tmp_path, vlm="# Т")
    ok = app._convert_uploaded_file(_FakeUpload("скан.pdf", _blank_pdf(1)), None, ocr_mode="mineru")
    broken = app._convert_uploaded_file(_FakeUpload("битый.pdf"), None, ocr_mode="mineru")
    assert ok["error"] is None and broken["error"] and broken["route"] is None
    for result in (ok, broken):                                 # и на успехе, и на ошибке
        assert set(result) == RESULT_KEYS


def test_app_no_longer_imports_run_pipeline():
    assert not hasattr(app, "run_pipeline")


# --- ПРАВКА #92: сверка по картинке ---------------------------------------------

import subprocess                                                   # noqa: E402

import ocr.vision                                                   # noqa: E402
from ocr import Finding                                             # noqa: E402
from ocr.vision import VISION_MODEL                                 # noqa: E402


@pytest.fixture
def no_claude(monkeypatch):
    """ПРАВКА #92: реальный claude не запускается ни в одном тесте блока."""
    def refuse(*args, **kwargs):
        raise AssertionError("тест запустил subprocess")
    monkeypatch.setattr(subprocess, "run", refuse)


_no_claude = pytest.mark.usefixtures("no_claude")


class _FakeStatus:
    def __init__(self):
        self.labels = []

    def update(self, label=None, state=None):
        self.labels.append(label)


def _fake_vision(monkeypatch):
    """Подмена ocr.vision.vision_findings: прогресс по двум страницам и одна vision_diff. Журнал вызовов."""
    calls: list[dict] = []

    def fake(md, pdf_bytes, zip_bytes, *, source_name, model=VISION_MODEL, progress=None):
        calls.append({"source_name": source_name, "model": model})
        if progress is not None:
            progress(1, 1, 2)
            progress(4, 2, 2)
        return [Finding(rule="vision_diff", severity="warning", page=1, snippet="Текст договора",
                        suggestion="Текст договоров", reading="… Текст договоров …", model=model)]

    monkeypatch.setattr(ocr.vision, "vision_findings", fake)
    return calls


def test_claude_available_follows_find_claude(monkeypatch):
    def missing():
        raise app.ClaudeCodeMissingError("нет")
    monkeypatch.setattr(app, "find_claude", missing)
    app._claude_available.clear()
    assert app._claude_available() is False
    monkeypatch.setattr(app, "find_claude", lambda: "C:/x/claude.cmd")
    app._claude_available.clear()
    assert app._claude_available() is True
    app._claude_available.clear()                       # не оставлять состояние другим тестам


def test_vision_progress_without_status_is_none():
    assert app._vision_progress(None) is None


def test_vision_rows():
    f = {"page": 3, "snippet": "IP65не", "reading": "… IP65 не ниже …", "suggestion": "IP65 не", "model": "m"}
    d = dict(f, snippet="лишнее", suggestion="")
    assert app._vision_rows([f, d]) == [
        {"Страница": 3, "Было": "IP65не", "По скану": "… IP65 не ниже …", "Стало": "IP65 не"},
        {"Страница": 3, "Было": "лишнее", "По скану": "… IP65 не ниже …", "Стало": "∅"}]


@_offline
@_no_claude
def test_mineru_mode_vision_passes_model_and_progress(monkeypatch, tmp_path):
    _fake_mineru(monkeypatch, tmp_path, vlm="# Договор\n\nТекст договора.")
    calls = _fake_vision(monkeypatch)
    status = _FakeStatus()

    result = app._convert_uploaded_file(_FakeUpload("скан.pdf", _blank_pdf(1)), None,
                                        ocr_mode="mineru", vision=True, status=status)

    assert result["error"] is None
    assert result["report"]["schema_version"] == 2
    vision = [f for f in result["report"]["findings"] if f["rule"].startswith("vision_")]
    assert len(vision) == 1 and vision[0]["rule"] == "vision_diff"
    assert vision[0]["reading"] and vision[0]["model"] == VISION_MODEL
    assert [c["model"] for c in calls] == [VISION_MODEL]
    progress = [label for label in status.labels if label and "из 2" in label]
    assert progress == ["Сверка по картинке: стр. 1 (1 из 2)…", "Сверка по картинке: стр. 4 (2 из 2)…"]
    assert set(result) == RESULT_KEYS


@_offline
@_no_claude
def test_mineru_mode_without_vision_does_not_call_it(monkeypatch, tmp_path):
    _fake_mineru(monkeypatch, tmp_path, vlm="# Договор\n\nТекст договора.")
    calls = _fake_vision(monkeypatch)

    result = app._convert_uploaded_file(_FakeUpload("скан.pdf", _blank_pdf(1)), None,
                                        ocr_mode="mineru", vision=False, status=_FakeStatus())

    assert result["error"] is None and calls == []
    assert not [f for f in result["report"]["findings"] if f["rule"].startswith("vision_")]


@_offline
@_no_claude
@pytest.mark.parametrize("name, data", [("док.docx", _docx_bytes()), ("табл.xlsx", _xlsx_bytes())],
                         ids=["docx", "xlsx"])
def test_vision_skipped_outside_mineru_routes(monkeypatch, tmp_path, name, data):
    _fake_mineru(monkeypatch, tmp_path, vlm="не должен понадобиться")
    calls = _fake_vision(monkeypatch)
    result = app._convert_uploaded_file(_FakeUpload(name, data), None, ocr_mode="mineru", vision=True)
    assert result["error"] is None                              # не ValueError
    assert result["route"] == "office" and calls == []


@_offline
@_no_claude
def test_vision_skipped_on_text_pdf(monkeypatch, tmp_path):
    pdf = require_fixture("textpdf1.pdf").read_bytes()          # стр. 1-2: слой есть, таблиц нет (спека 09)
    _fake_mineru(monkeypatch, tmp_path, vlm="x")
    calls = _fake_vision(monkeypatch)
    result = app._convert_uploaded_file(_FakeUpload("т.pdf", pdf), "1-2", ocr_mode="mineru", vision=True)
    assert result["error"] is None
    assert result["route"] == "text" and calls == []


@_offline
@_no_claude
def test_vision_failure_keeps_conversion(monkeypatch, tmp_path):
    _fake_mineru(monkeypatch, tmp_path, vlm="# Договор\n\nТекст договора.")   # zip только с full.md
    made: list[str] = []

    def broken_verifier(model):
        made.append(model)
        raise AssertionError("верификатор не должен создаваться без content_list")
    monkeypatch.setattr(ocr.vision, "make_verifier", broken_verifier)

    result = app._convert_uploaded_file(_FakeUpload("скан.pdf", _blank_pdf(1)), None,
                                        ocr_mode="mineru", vision=True, status=_FakeStatus())

    assert result["error"] is None and "Текст договора." in result["markdown"]
    skipped = [f for f in result["report"]["findings"] if f["rule"] == "vision_skipped"]
    assert len(skipped) == 1
    assert skipped[0]["suggestion"].startswith("сверка не выполнена: ValueError: ")
    assert made == []
