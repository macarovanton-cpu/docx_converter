"""Тесты табличного рендерера convert.py (запускать из docx_converter/)."""

import pytest
from docx import Document
from docx.oxml.ns import qn

from convert import convert_md_to_docx

# ПРАВКА #34: иконка добавляется к тексту, а не вместо него
CASES = [
    ("Да — 36 месяцев",       "✓ Да — 36 месяцев"),
    ("Нет данных",            "✗ Нет данных"),
    ("Нет.",                  "✗ Нет."),
    ("Да",                    "✓ Да"),
    ("Отсутствует в базовой", "✗ Отсутствует в базовой"),
    ("Данные по объекту",     "Данные по объекту"),   # «Да» без \b — иконки нет
    ("Гарантия",              "Гарантия"),
]


@pytest.fixture(scope="module")
def table(tmp_path_factory):
    """Одна конвертация без шаблона на весь модуль → таблица из готового docx."""
    rows = "\n".join(f"| Параметр {i} | {src} |" for i, (src, _) in enumerate(CASES))
    md = f"| Параметр | Значение |\n|---|---|\n{rows}\n"
    out = tmp_path_factory.mktemp("docx") / "table.docx"
    convert_md_to_docx(md, str(out))
    return Document(str(out)).tables[0]


@pytest.fixture(scope="module")
def cells(table):
    """{исходный текст ячейки: то, что реально попало в word/document.xml}."""
    return {src: row.cells[1].text for (src, _), row in zip(CASES, table.rows[1:])}


def test_table_row_count(table):
    """Страховка: шапка + все строки на месте, иначе ассерты ниже вакуумны."""
    assert len(table.rows) == len(CASES) + 1


@pytest.mark.parametrize("src,expected", CASES)
def test_cell_text_survives_icon(cells, src, expected):
    assert cells[src] == expected


def test_header_repeats_and_rows_do_not_split(table):
    """ПРАВКА #35: tblHeader только на шапке, cantSplit на каждой строке."""
    headers = [i for i, r in enumerate(table.rows)
               if r._tr.find(qn('w:trPr')).find(qn('w:tblHeader')) is not None]
    assert headers == [0]
    assert all(r._tr.find(qn('w:trPr')).find(qn('w:cantSplit')) is not None
               for r in table.rows)


# =============================================================================
# ИНЛАЙН-ПАРСЕР
# =============================================================================

def _runs(md, tmp_path):
    """[(текст, bold, italic), ...] по всем абзацам конвертации без шаблона."""
    out = tmp_path / "inline.docx"
    convert_md_to_docx(md, str(out))
    return [(r.text, bool(r.bold), bool(r.italic))
            for p in Document(str(out)).paragraphs for r in p.runs]


def _text(md, tmp_path):
    return "".join(t for t, _, _ in _runs(md, tmp_path))


# ПРАВКА #36: звёздочки-умножение не съедаются как курсив
def test_multiplication_with_spaces_survives(tmp_path):
    runs = _runs("Габариты платформы 2 * 3 * 4 метра.", tmp_path)
    assert "2 * 3 * 4" in "".join(t for t, _, _ in runs)
    assert not any(italic for _, _, italic in runs)


def test_multiplication_without_spaces_survives(tmp_path):
    runs = _runs("Габариты 2*3*4 метра.", tmp_path)
    assert "2*3*4" in "".join(t for t, _, _ in runs)
    assert not any(italic for _, _, italic in runs)


def test_italic_still_works(tmp_path):
    assert ("курсив", False, True) in _runs("Обычный *курсив* тут.", tmp_path)


def test_bold_still_works(tmp_path):
    assert ("жирный", True, False) in _runs("И **жирный** тут.", tmp_path)


def test_bold_italic_still_works(tmp_path):
    assert ("оба", True, True) in _runs("И ***оба*** тут.", tmp_path)


def test_emphasis_with_inner_space_still_works(tmp_path):
    """Пробелы внутри эмфазы разрешены — правило касается только флангов."""
    assert ("две слова", False, True) in _runs("Тут *две слова* курсивом.",
                                               tmp_path)


# ПРАВКА #36: +? вместо *? — «****» не матчится с пустым содержимым и не пропадает
def test_four_asterisks_survive_literally(tmp_path):
    assert "****" in _text("Тут **** четыре звезды.", tmp_path)


# ПРАВКА #37: экранированные спецсимволы выводятся буквально
def test_escaped_asterisks_are_literal(tmp_path):
    runs = _runs(r"Звёздочки: \*звёздочки вокруг слова\* — курсива нет.",
                 tmp_path)
    assert "*звёздочки вокруг слова*" in "".join(t for t, _, _ in runs)
    assert not any(italic for _, _, italic in runs)


def test_escaped_specials_lose_backslash(tmp_path):
    md = (r"Решётка \# и скобки \[ГОСТ 29329-92\], черта \| и черта \_, "
          r"кавычка \` и стрелка \> — все на месте.")
    text = _text(md, tmp_path)
    assert "\\" not in text
    for symbol in "#[]|_`>":
        assert symbol in text, symbol


def test_missing_image_caption_unshields_escaped_brackets(tmp_path):
    md = r"Текст перед ![Схема \[черновик\]](missing.png) текст после."
    text = _text(md, tmp_path)
    assert "[черновик]" in text
    assert not any(0xE000 <= ord(c) <= 0xE0FF for c in text)


def test_escaped_brackets_do_not_create_link(tmp_path):
    """Экранированные скобки не должны собраться в гиперссылку с (url)."""
    out = tmp_path / "esc_link.docx"
    convert_md_to_docx(r"Ссылка \[ГОСТ 29329-92\](https://example.com) нет.",
                       str(out))
    doc = Document(str(out))
    assert "[ГОСТ 29329-92](https://example.com)" in "".join(
        p.text for p in doc.paragraphs)
    assert not doc.element.body.findall('.//' + qn('w:hyperlink'))


def test_real_link_still_works(tmp_path):
    out = tmp_path / "link.docx"
    convert_md_to_docx("Ссылка [сайт](https://tenzosila.ru) работает.",
                       str(out))
    doc = Document(str(out))
    assert doc.element.body.findall('.//' + qn('w:hyperlink'))


def test_no_placeholder_leaks(tmp_path):
    """Ни один символ private use area не должен доехать до документа."""
    md = (r"Звёздочка \*, скобки \[x\], ссылка [сайт](https://tenzosila.ru), "
          r"курсив *да* и жирный **да**.")
    assert not [c for c in _text(md, tmp_path)
            if chr(0xE000) <= c <= chr(0xF8FF)]
