"""Оформление письма (ПРАВКИ #55-#58): строгий бланк вместо брендового стиля ПЗ.

Регрессию ПЗ ловит tests/test_golden_pz.py — здесь проверяется только то, чем
письмо отличается, и то, что пре-проход шапки ничего не теряет.
"""

import zipfile

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

from convert import STYLE_LETTER, STYLE_PZ, convert_md_to_docx

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'

LETTER = """**Дата:** 18.09.2026
**Исх.:** № 145/26

# ПИСЬМО

*О цене работ по Договору №12/25*

**Кому:** ООО «Заказчик»
Генеральному директору И.И. Петрову

## 1. О цене работ по Договору

Текст раздела со [ссылкой](https://tenzosila.ru).

Второй абзац раздела — он идёт не сразу за заголовком, поэтому красная строка
к нему применяется.

- первый пункт
- второй пункт

**Кому:** Второй адресат в теле

С уважением,
директор ООО «ТПК «Тензосила»
Сенаторов О.А.
"""


def _build(md, tmp_path, name="letter.docx", doc_style="letter"):
    out = tmp_path / name
    convert_md_to_docx(md, str(out), doc_style=doc_style)
    return out


@pytest.fixture(scope="module")
def letter(tmp_path_factory):
    return _build(LETTER, tmp_path_factory.mktemp("letter"))


@pytest.fixture(scope="module")
def doc(letter):
    return Document(str(letter))


def _pPr(paragraph):
    return paragraph._p.find(qn('w:pPr'))


def _align(paragraph):
    pPr = _pPr(paragraph)
    jc = None if pPr is None else pPr.find(qn('w:jc'))
    return None if jc is None else jc.get(qn('w:val'))


def _border(paragraph, side):
    pPr = _pPr(paragraph)
    pBdr = None if pPr is None else pPr.find(qn('w:pBdr'))
    return None if pBdr is None else pBdr.find(qn(f'w:{side}'))


def _find(doc, text):
    """Первый абзац тела, начинающийся с text."""
    return next(p for p in doc.paragraphs if p.text.strip().startswith(text))


# =============================================================================
# ПРАВКА #56: заголовок, тема, реквизиты, подпись
# =============================================================================

def test_h1_is_centered_and_black(doc):
    p = _find(doc, "ПИСЬМО")
    assert _align(p) == "center"
    assert all(str(r.font.color.rgb) == "000000" for r in p.runs)


def test_no_decorative_rule_anywhere_in_letter(doc):
    """Красной линии под H1 в бланке нет — и вообще никакой нижней границы
    абзаца в теле, кроме линии под шапкой (она внутри таблицы)."""
    assert all(_border(p, 'bottom') is None for p in doc.paragraphs)


def test_subject_line_is_centered_italic(doc):
    p = _find(doc, "О цене работ по Договору №12/25")
    assert _align(p) == "center"
    assert all(r.italic for r in p.runs)


def test_subject_is_not_wrapped_in_intro_band(doc):
    """В ПЗ тот же блок становится таблицей-врезкой, в письме — абзацем."""
    assert "О цене работ" not in "".join(c.text for t in doc.tables
                                         for r in t.rows for c in r.cells)


def test_section_heading_is_body_size_and_black(doc):
    p = _find(doc, "1. О цене работ по Договору")
    assert all(r.font.size == Pt(STYLE_LETTER['h2_size']) for r in p.runs)
    assert all(str(r.font.color.rgb) == "000000" for r in p.runs)


def test_requisites_in_body_are_right_aligned_without_fill(doc):
    p = _find(doc, "Кому: Второй адресат")
    assert _align(p) == "right"
    assert _pPr(p).find(qn('w:shd')) is None


def test_signature_has_no_rule_and_keeps_left_alignment(doc):
    """Правый tab stop не работает на абзаце, выключенном вправо."""
    p = _find(doc, "С уважением")
    assert _border(p, 'top') is None
    assert _align(p) == "left"


def test_signature_position_and_name_share_one_line(doc):
    p = _find(doc, "С уважением")
    tabs = _pPr(p).find(qn('w:tabs'))
    assert tabs is not None, "нет табуляции в подписи"
    assert tabs.find(qn('w:tab')).get(qn('w:val')) == "right"
    assert p._p.findall('.//' + qn('w:tab')), "нет символа табуляции между должностью и фамилией"


def test_short_signature_is_not_glued(tmp_path):
    """Две строки — склеивать нечего, табуляции не появляется."""
    doc = Document(str(_build("С уважением,\nСенаторов О.А.", tmp_path, "short.docx")))
    p = _find(doc, "С уважением")
    assert _pPr(p).find(qn('w:tabs')) is None


def test_link_is_black_and_underlined(letter):
    """Без цвета подчёркивание — единственный признак ссылки."""
    with zipfile.ZipFile(str(letter)) as z:
        xml = z.read('word/document.xml').decode('utf-8')
    hyperlink = xml[xml.index('<w:hyperlink'):xml.index('</w:hyperlink>')]
    assert '<w:color w:val="000000"/>' in hyperlink
    assert '<w:u w:val="single"/>' in hyperlink


def test_body_size_is_smaller_than_pz(doc):
    assert doc.styles['Normal'].font.size == Pt(STYLE_LETTER['body_size'])
    assert STYLE_LETTER['body_size'] < STYLE_PZ['body_size']


def test_first_line_indent_is_zero(doc, tmp_path_factory):
    """В деловом бланке красной строки нет, в ПЗ она осталась 0.75 см."""
    # сравнение с допуском: в XML отступ лежит в twips, и 0.75 см проходит
    # round-trip как 425 twips = 269875 EMU, а не ровно 270000
    p = _find(doc, "Второй абзац раздела")
    assert p.paragraph_format.first_line_indent.cm == pytest.approx(
        STYLE_LETTER['para_indent'], abs=0.01)
    pz = Document(str(_build(LETTER, tmp_path_factory.mktemp("pz"),
                             "pz.docx", doc_style="pz")))
    assert (_find(pz, "Второй абзац раздела").paragraph_format.first_line_indent.cm
            == pytest.approx(STYLE_PZ['para_indent'], abs=0.01))


# =============================================================================
# ПРАВКА #57: шапка «Дата/Исх. + адресат»
# =============================================================================

def test_header_is_first_in_body(doc):
    first = next(el.tag.split('}')[-1] for el in doc.element.body.iterchildren())
    assert first == "tbl"


def test_header_is_one_row_two_borderless_cells(doc):
    cells = doc.tables[0].rows[0].cells
    assert len(doc.tables[0].rows) == 1 and len(cells) == 2
    for cell in cells:
        borders = cell._tc.find(qn('w:tcPr')).find(qn('w:tcBorders'))
        assert all(b.get(qn('w:val')) == "none" for b in borders)


def test_header_left_column_width_is_fixed(doc):
    """Без tcW и autofit=False Word разложит колонки по содержимому."""
    left = doc.tables[0].rows[0].cells[0]
    tcW = left._tc.find(qn('w:tcPr')).find(qn('w:tcW'))
    assert int(tcW.get(qn('w:w'))) == int(STYLE_LETTER['header_left_cm'] * 567)


def test_rule_spans_only_the_left_column(doc):
    """Линия — граница абзаца внутри левой ячейки, а не таблицы или страницы."""
    left, right = doc.tables[0].rows[0].cells
    assert _border(left.paragraphs[0], 'bottom') is not None
    assert _border(right.paragraphs[0], 'bottom') is None


def test_header_holds_date_and_addressee(doc):
    left, right = doc.tables[0].rows[0].cells
    assert "18.09.2026" in left.text and "145/26" in left.text
    assert "Заказчик" in right.text or "адресат" in right.text.lower()


def test_addressee_label_is_stripped_and_right_aligned(tmp_path):
    doc = Document(str(_build(
        "**Кому:** ООО «Заказчик»\n\n# ПИСЬМО\n\nТекст.", tmp_path, "komu.docx")))
    right = doc.tables[0].rows[0].cells[1]
    assert right.text.strip() == "ООО «Заказчик»"
    assert _align(right.paragraphs[0]) == "right"


# =============================================================================
# ПРАВКА #57: пре-проход не теряет блоки (главный риск конструкции)
# =============================================================================

@pytest.mark.parametrize("md,expect_in_header,expect_in_body", [
    ("**Дата:** 18.09.2026\n\n# ПИСЬМО\n\nТекст.", "18.09.2026", "Текст."),
    ("**Кому:** Заказчик\n\n# ПИСЬМО\n\nТекст.", "Заказчик", "Текст."),
    ("# ПИСЬМО\n\nТекст.", None, "Текст."),
])
def test_header_without_one_half_loses_nothing(md, expect_in_header,
                                               expect_in_body, tmp_path):
    doc = Document(str(_build(md, tmp_path, f"half{len(md)}.docx")))
    body = " ".join(p.text for p in doc.paragraphs)
    header = " ".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert expect_in_body in body
    if expect_in_header is None:
        assert not doc.tables, "шапки быть не должно — оба блока отсутствуют"
    else:
        assert expect_in_header in header


def test_second_requisites_block_stays_in_body(doc):
    """В шапку уходит только первый **Кому:**, второй остаётся абзацем."""
    assert "Второй адресат" in " ".join(p.text for p in doc.paragraphs)


def test_meta_block_in_pz_is_not_swallowed(tmp_path):
    """В ПЗ пре-прохода нет: блок **Дата:** рендерится обычным абзацем."""
    doc = Document(str(_build("**Дата:** 18.09.2026\n\n# ПЗ\n\nТекст.",
                              tmp_path, "pzmeta.docx", doc_style="pz")))
    assert any("18.09.2026" in p.text for p in doc.paragraphs)


# =============================================================================
# ПРАВКА #58: маркер списка
# =============================================================================

@pytest.mark.parametrize("doc_style,glyph,font", [
    ("letter", "—", "PT Sans"),
    ("pz", "•", "Symbol"),
])
def test_bullet_glyph_and_font_come_from_profile(doc_style, glyph, font, tmp_path):
    """В Symbol U+2014 — не тире, поэтому шрифт обязан ехать вместе с глифом."""
    from lxml import etree
    out = _build("- пункт", tmp_path, f"bul_{doc_style}.docx", doc_style)
    with zipfile.ZipFile(str(out)) as z:
        root = etree.fromstring(z.read('word/numbering.xml'))
    abstract = next(a for a in root.findall(f'{{{W}}}abstractNum')
                    if a.get(f'{{{W}}}abstractNumId') == '100')
    lvl = abstract.find(f'{{{W}}}lvl')
    assert lvl.find(f'{{{W}}}lvlText').get(f'{{{W}}}val') == glyph
    rFonts = lvl.find(f'{{{W}}}rPr').find(f'{{{W}}}rFonts')
    assert rFonts.get(f'{{{W}}}ascii') == font


# =============================================================================
# ПРАВКА #55: сам профиль
# =============================================================================

def test_profiles_have_identical_key_sets():
    """Ключ, забытый в одном профиле, — KeyError посреди рендера."""
    assert STYLE_PZ.keys() == STYLE_LETTER.keys()


def test_unknown_doc_style_fails_loudly(tmp_path):
    """Молча отрисованное не тем стилем письмо хуже ошибки в UI."""
    with pytest.raises(KeyError):
        convert_md_to_docx("# Тест", str(tmp_path / "x.docx"), doc_style="опечатка")


def test_letter_uses_no_brand_colors(letter):
    with zipfile.ZipFile(str(letter)) as z:
        xml = z.read('word/document.xml').decode('utf-8')
    for brand in ("015198", "D04514", "EF7F1A", "EBF3FB", "FFF8F0", "1A1A1A"):
        assert brand not in xml, f"в письме остался фирменный цвет {brand}"
